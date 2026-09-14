# -*- coding: utf-8 -*-
"""DOCX eksport: Shablons/Ariza/*.docx fayllarni to'ldirib, tayyor hujjat qaytaradi.

Run (bir necha <w:r> ga bo'lingan) chegaralaridan qat'iy nazar matn almashtirish
mantig'i core/docx_utils.py da — u "Reestr tizimi" bilan umumiy.
"""
import io
from html import escape

from django.conf import settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .docx_utils import ShablonTopilmadi, expand_paragraph, find_paragraph, replace_in_doc
from .reasons import TEMPLATE_FAYLLAR
from .text_utils import ijrochi_qisqa_ism

PLACEHOLDER = "{Appda belgilangan sabablar}"

OYLAR = {
    1: "yanvar", 2: "fevral", 3: "mart", 4: "aprel",
    5: "may", 6: "iyun", 7: "iyul", 8: "avgust",
    9: "sentyabr", 10: "oktyabr", 11: "noyabr", 12: "dekabr",
}


def format_sana(d):
    if not d:
        return ""
    return f"{d.year}-yil {d.day}-{OYLAR[d.month]}"


def shablon_yolini_topish(kategoriya, holat):
    """(kategoriya, holat) juftligi uchun shablon faylini qaytaradi.

    Ilgari bu yer to'g'ridan-to'g'ri TEMPLATE_FAYLLAR[...] edi: kategoriya
    ro'yxatdan chiqarilgan bo'lsa (masalan "favqulodda") yoki shablon fayli
    o'chib ketgan bo'lsa, foydalanuvchi 500-xatolik ko'rardi. Endi sababi
    tushunarli xabar beriladi.
    """
    fayl_nomi = TEMPLATE_FAYLLAR.get((kategoriya, holat))
    if not fayl_nomi:
        raise ShablonTopilmadi(
            f"\"{kategoriya}\" kategoriyasi va \"{holat}\" holati uchun shablon "
            "belgilanmagan. Bu kategoriya tizimdan olib tashlangan bo'lishi mumkin — "
            "arizani mavjud kategoriyalardan biriga o'zgartiring."
        )
    yol = settings.SHABLONLAR_DIR / fayl_nomi
    if not yol.exists():
        raise ShablonTopilmadi(
            f"Shablon fayli topilmadi: Shablons/Ariza/{fayl_nomi}"
        )
    return yol, fayl_nomi


def _imzo_malumotlari(ariza):
    """Direktor/ijrochi ma'lumotlari — .docx to'ldirishda ham, preview'da ham
    xuddi shu joydan olinadi, shuning uchun ikkalasi hech qachon farq qilmaydi."""
    profil = getattr(ariza.created_by, "profil", None)
    tashkilot = profil.tashkilot if profil else None
    tashkilot_nomi = tashkilot.nomi if tashkilot else "Tashkilot"
    tashkilot_rahbar = tashkilot.rahbar if tashkilot else ""
    ijrochi = ijrochi_qisqa_ism(ariza.created_by)
    return tashkilot_nomi, tashkilot_rahbar, ijrochi


def ariza_docx_yaratish(ariza):
    """Ariza obyekti asosida to'ldirilgan docx faylni BytesIO sifatida qaytaradi."""
    shablon_yoli, fayl_nomi = shablon_yolini_topish(ariza.kategoriya, ariza.holat)

    doc = Document(str(shablon_yoli))

    tashkilot_nomi, tashkilot_rahbar, ijrochi = _imzo_malumotlari(ariza)

    replacements = {
        "{tuman}": ariza.tuman,
        "{mfy}": ariza.mfy,
        "{kucha}": ariza.kucha,
        "{fio}": ariza.fio,
        "{sana}": format_sana(ariza.sana),
        "{murojaat_raqami}": ariza.murojaat_raqami,
        "{ariza_raqami}": ariza.ariza_raqami,
        "{tashkilot}": ariza.tashkilot,
        # Shablonlarda statik yozilgan imzo bloki — endi tashkilot va
        # xodimning o'ziga moslab almashtiriladi.
        "Markazi direktori:": f"{tashkilot_nomi} direktori:",
        "S.Mutalibov": tashkilot_rahbar,
        "Ijrochi: Y.Olimjonov": f"Ijrochi: {ijrochi}",
    }

    if ariza.holat == "tayinlangan":
        replacements["{ajratilgan_summa}"] = ariza.ajratilgan_summa
        replacements["{kollegal_qaror}"] = ariza.kollegal_qaror
        replace_in_doc(doc, replacements)
    else:
        replace_in_doc(doc, replacements)
        sabablar = ariza.rad_sabablari_royxati()
        matnlar = [f"{i + 1}) {t}" for i, (_, t) in enumerate(sabablar)] or ["Sabab ko'rsatilmagan."]
        marker = find_paragraph(doc, PLACEHOLDER)
        if marker is not None:
            expand_paragraph(marker, matnlar)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer, fayl_nomi


def _runlarni_html_qilish(paragraph):
    """Paragraf runlarini (bold/italic holatini saqlab) HTML qatoriga yig'adi.

    Ketma-ket kelgan bir xil formatdagi runlar birlashtiriladi — Word docx
    matnni ko'plab kichik runlarga bo'lib saqlaydi (masalan bitta so'z ham
    bir necha run bo'lishi mumkin), shuning uchun run boshiga qarab emas,
    holat o'zgarganda ajratib chiqiladi.
    """
    qismlar = []
    joriy_matn = []
    joriy_holat = None
    for run in paragraph.runs:
        matn = run.text.replace("\xa0", " ")
        if not matn:
            continue
        holat = (bool(run.bold), bool(run.italic))
        if joriy_matn and holat != joriy_holat:
            qismlar.append((joriy_holat, "".join(joriy_matn)))
            joriy_matn = []
        joriy_holat = holat
        joriy_matn.append(matn)
    if joriy_matn:
        qismlar.append((joriy_holat, "".join(joriy_matn)))

    html_boo = []
    for (bold, italic), matn in qismlar:
        parcha = escape(matn)
        if italic:
            parcha = f"<i>{parcha}</i>"
        if bold:
            parcha = f"<b>{parcha}</b>"
        html_boo.append(parcha)
    return "".join(html_boo)


def ariza_preview_malumotlari(ariza):
    """Ariza uchun HTML preview'da ko'rsatiladigan ma'lumotlarni qaytaradi.

    Muhim: matn to'g'ridan-to'g'ri eksport qilinadigan (haqiqiy) .docx dan
    o'qiladi — alohida "preview matni" moduli yozilmagan. Shu sababli preview
    va yuklab olinadigan hujjat matni hech qachon bir-biridan farq qilmaydi
    (Reestr tizimida bunday ajralib qolish bir necha marta xatolikka olib
    kelgan edi).

    Imzo qatori (direktor/ijrochi) alohida ajratib olinadi va toza HTML bilan
    chiziladi — chunki shablon faylida bu qator Wordda tekislash uchun juda
    ko'p \xa0/tab belgilari bilan yozilgan va uni o'zgarishsiz veb-sahifaga
    chiqarish chiroyli ko'rinmaydi. Qiymatlar (tashkilot nomi, rahbar, ijrochi)
    esa aynan .docx ni to'ldirishda ishlatilgan joydan olinadi.
    """
    buffer, _fayl_nomi = ariza_docx_yaratish(ariza)
    doc = Document(buffer)
    tashkilot_nomi, tashkilot_rahbar, ijrochi = _imzo_malumotlari(ariza)

    paragraphs = []
    for p in doc.paragraphs:
        oddiy_matn = p.text.replace("\xa0", " ").strip()
        if not oddiy_matn:
            continue
        if "direktori:" in oddiy_matn or oddiy_matn.startswith("Ijrochi:"):
            continue
        paragraphs.append({
            "html": _runlarni_html_qilish(p),
            "center": p.alignment == WD_ALIGN_PARAGRAPH.CENTER,
        })

    return {
        "paragraphs": paragraphs,
        "tashkilot_nomi": tashkilot_nomi,
        "tashkilot_rahbar": tashkilot_rahbar,
        "ijrochi": ijrochi,
    }
